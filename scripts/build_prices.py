#!/usr/bin/env python3
"""
Builds small per-country fuel price files for the Tankpot app from the
European Commission's Weekly Oil Bulletin (prices with taxes, all EU countries).

Output: eu/<CC>.json   e.g. eu/FR.json
  {
    "country": "FR",
    "source": "European Commission, Weekly Oil Bulletin (prices with taxes)",
    "unit": "EUR/L",
    "generated": "2026-10-12T07:23:00Z",
    "columns": ["date", "euro95", "diesel", "lpg"],
    "weeks": [["2026-10-05", 1.812, 1.701, 0.99], ...]   # newest last, date = prices in force on
  }
plus eu/index.json listing countries and the newest week.

Usage: build_prices.py [path-to-history.xlsx]   (downloads it when no path is given)
"""
import datetime as dt
import json
import os
import re
import sys
import urllib.request

import openpyxl

PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
# Known location of the history file; the page is searched first in case it moves.
FALLBACK_URL = ("https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
                "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "eu")
KEEP_WEEKS = 160  # about three years
UA = {"User-Agent": "tankpot-prices/1.0 (+https://github.com/HoundRock/tankpot-prices)"}

PRODUCTS = {  # header keyword -> output column
    "euro95": "euro95", "euro_95": "euro95", "eurosuper95": "euro95", "euro-super95": "euro95",
    "diesel": "diesel", "gasoil": "diesel", "gas_oil": "diesel", "automotivegasoil": "diesel",
    "lpg": "lpg",
}
COLUMNS = ["euro95", "diesel", "lpg"]


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()


def find_history_url() -> str:
    try:
        html = fetch(PAGE).decode("utf-8", "replace")
        links = re.findall(r'href="([^"]+)"', html)
        for link in links:
            if "history" in link.lower() and ("xlsx" in link.lower() or "/document/download/" in link):
                return link if link.startswith("http") else "https://energy.ec.europa.eu" + link
    except Exception as e:  # page layout changes shouldn't break the job
        print("Could not read bulletin page:", e)
    return FALLBACK_URL


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def as_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def parse_header(cell):
    """'FR_price_with_tax_euro95' -> ('FR', 'euro95'); returns None for anything else."""
    if not isinstance(cell, str):
        return None
    m = re.match(r"^\s*([A-Za-z]{2})[_\s-]+price[_\s-]+with[_\s-]+tax(?:es)?[_\s-]+(.+?)\s*$", cell)
    if not m:
        return None
    product = re.sub(r"[\s]", "", m.group(2).lower())
    country = m.group(1).upper()
    country = {"EL": "GR", "UK": "GB"}.get(country, country)  # EU codes -> ISO codes used by phones
    for key, col in PRODUCTS.items():
        if product.startswith(key):
            return country, col
    return None


def parse(path: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {}  # country -> {date -> {col: price}}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header_map = None
        for row in rows:
            if header_map is None:
                found = {i: parse_header(c) for i, c in enumerate(row)}
                found = {i: v for i, v in found.items() if v}
                if len(found) >= 10:
                    header_map = found
                    print(f"Sheet '{ws.title}': {len(found)} price columns")
                continue
            date = next((as_date(c) for c in row[:3] if as_date(c)), None)
            if date is None:
                continue
            for i, (cc, col) in header_map.items():
                if i >= len(row):
                    continue
                val = as_number(row[i])
                if val is None or val <= 0:
                    continue
                per_litre = val / 1000.0 if val > 20 else val  # bulletin quotes EUR per 1000 L
                data.setdefault(cc, {}).setdefault(date, {})[col] = round(per_litre, 4)
        if data:
            break  # first sheet with prices is "with taxes"
    if not data:
        # Print the first rows so a format change is easy to diagnose in the job log.
        for ws in wb.worksheets:
            print("== sheet", ws.title)
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                print([c for c in row[:12]])
                if n > 6:
                    break
        raise SystemExit("No price columns found; the bulletin format may have changed.")
    return data


def write(data):
    os.makedirs(OUT_DIR, exist_ok=True)
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    index = {"generated": now, "countries": {}}
    for cc, by_date in sorted(data.items()):
        dates = sorted(by_date)[-KEEP_WEEKS:]
        weeks = [[d.isoformat()] + [by_date[d].get(c) for c in COLUMNS] for d in dates]
        doc = {
            "country": cc,
            "source": "European Commission, Weekly Oil Bulletin (prices with taxes)",
            "unit": "EUR/L",
            "generated": now,
            "columns": ["date"] + COLUMNS,
            "weeks": weeks,
        }
        with open(os.path.join(OUT_DIR, f"{cc}.json"), "w") as f:
            json.dump(doc, f, separators=(",", ":"))
        index["countries"][cc] = weeks[-1][0] if weeks else None
    with open(os.path.join(OUT_DIR, "index.json"), "w") as f:
        json.dump(index, f, indent=1, sort_keys=True)
    print(f"Wrote {len(data)} countries; newest week:", max(v for v in index["countries"].values() if v))


def main():
    if len(sys.argv) > 1:
        path = sys.argv[1]
    else:
        url = find_history_url()
        print("Downloading", url)
        path = "history.xlsx"
        with open(path, "wb") as f:
            f.write(fetch(url))
    write(parse(path))


if __name__ == "__main__":
    main()
